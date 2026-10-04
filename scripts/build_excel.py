from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from src.model import build, PROC, OUT
from src.scenarios import scenario

ROOT=Path(__file__).resolve().parents[1]
if not (PROC/'fpna_model.csv').exists(): build(use_demo=True)
m=pd.read_csv(PROC/'fpna_model.csv',parse_dates=['month']); pnl=pd.read_csv(OUT/'pnl_monthly.csv',parse_dates=['month'])
path=ROOT/'excel/FPNA_Enterprise_Model.xlsx'
wb=Workbook(); ws=wb.active; ws.title='Executive'
ws['A1']='Enterprise FP&A Performance Model'; ws['A1'].font=Font(size=18,bold=True)
ws['A3']='Metric'; ws['B3']='Actual'; ws['C3']='Budget'; ws['D3']='Variance'; ws['E3']='Variance %'
for c in ws[3]: c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor='D9E2F3')
agg=m.groupby('category')[['actual','budget']].sum(); order=['Revenue','COGS','Personnel','Opex','Depreciation']
r=4
for cat in order:
    a=float(agg.loc[cat,'actual']) if cat in agg.index else 0; b=float(agg.loc[cat,'budget']) if cat in agg.index else 0
    ws.cell(r,1,cat); ws.cell(r,2,a); ws.cell(r,3,b); ws.cell(r,4,a-b); ws.cell(r,5,(a-b)/b if b else None); r+=1
ws.cell(r,1,'Gross Profit'); ws.cell(r,2,'=B4+B5'); ws.cell(r,3,'=C4+C5'); ws.cell(r,4,'=B9-C9'); ws.cell(r,5,'=D9/C9')
for name,df in [('Model',m),('Monthly P&L',pnl)]:
    sh=wb.create_sheet(name); sh.append(list(df.columns))
    for row in df.itertuples(index=False): sh.append(list(row))
sc=wb.create_sheet('Scenarios'); sc.append(['Scenario','Revenue Growth','COGS Change','Personnel Growth','Opex Growth','Operating Profit','Operating Margin'])
base=scenario(pnl); cases=[('Base',0,0,0,0),('Upside',.05,0,0,0),('Downside',-.05,.02,.03,.03)]
for name,rg,cg,pg,og in cases:
    s=scenario(pnl,rg,cg,pg,og); sc.append([name,rg,cg,pg,og,s['Operating Profit'],s['Operating Margin']])
wb.save(path); print(path)
